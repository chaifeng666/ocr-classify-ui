#!/usr/bin/env python3
"""產生 SBOM（軟體物料清單）：sbom/sbom.cdx.json（CycloneDX 1.5）、sbom/SBOM.md、SBOM.xlsx。

涵蓋兩個應用：
  A. ocr-classify-ui（本 repo，GitHub Pages 網路版＋離線版 ZIP）
  B. ocr-app（ocr.html＋OCR-offline.zip，本地 OCR 小程式）

用法：python3 tools/sbom.py [OCR-offline 解壓縮資料夾] [xlsx 輸出路徑]
"""
import hashlib, json, os, sys, uuid, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OCR_DIR = sys.argv[1] if len(sys.argv) > 1 else None
XLSX_OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'sbom', 'SBOM.xlsx')

OK, WARN, GONE = '✅ 通過', '⚠️ 需注意', '🗑️ 已移除'

def sha(path):
    if not path or not os.path.exists(path): return ''
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(1 << 20), b''): h.update(b)
    return h.hexdigest()

def ocr(rel): return os.path.join(OCR_DIR, rel) if OCR_DIR else None

# 每個元件：應用、名稱、版本、類型、授權、開發者/組織、來源國、取得方式、用途、檔案、結論、說明
C = []
def add(app, name, version, ctype, lic, supplier, country, source, use, result, note='', path=None, purl=''):
    C.append(dict(app=app, name=name, version=version, type=ctype, license=lic, supplier=supplier, country=country,
                  source=source, use=use, result=result, note=note, sha256=sha(path), purl=purl))

A, B = 'ocr-classify-ui（單據分類，GitHub Pages／離線版）', 'ocr-app（本地 OCR 小程式 ocr.html／OCR-offline.zip）'

# ---------- A. ocr-classify-ui ----------
add(A, 'ocr-classify-ui index.html（自行開發）', 'main', 'application', '專案自有', 'chaifeng666', '臺灣',
    'GitHub repo chaifeng666/ocr-classify-ui', '主程式（原生 JavaScript，無框架）', OK, '不含任何第三方程式碼',
    os.path.join(ROOT, 'index.html'))
add(A, 'SheetJS（xlsx）', '0.18.5', 'library', 'Apache-2.0', 'SheetJS LLC', '美國',
    'npm registry → 放在 repo lib/（自行託管，不走 CDN）', '匯出 Excel', OK, '',
    os.path.join(ROOT, 'lib', 'xlsx.full.min.js'), 'pkg:npm/xlsx@0.18.5')
for nm, ver, lic, sup, src, use, rel, purl in [
    ('tesseract.js（含 worker）', '5.1.1', 'Apache-2.0', 'Project Naptha（Kevin Kwok 等）', 'npm registry → repo lib/', 'OCR 主程式', 'lib/tesseract.min.js', 'pkg:npm/tesseract.js@5.1.1'),
    ('tesseract.js-core（Tesseract OCR 5＋Leptonica，WebAssembly）', '5.1.1', 'Apache-2.0（Tesseract）／BSD-2-Clause（Leptonica）', 'Tesseract：HP → Google；Leptonica：Dan Bloomberg', 'npm registry → repo lib/core/', 'OCR 引擎', 'lib/core/tesseract-core-simd-lstm.wasm.js', 'pkg:npm/tesseract.js-core@5.1.1'),
    ('Tesseract 語言資料 chi_tra＋eng', '4.0.0 best_int', 'Apache-2.0', 'Google LLC（tessdata）', 'repo lib/lang/', '繁中／英文辨識模型', 'lib/lang/chi_tra.traineddata.gz', 'pkg:generic/tessdata/chi_tra@4.0.0'),
    ('PDF.js（pdfjs-dist，含 worker 與 Adobe CMap）', '3.11.174', 'Apache-2.0（CMap：BSD-3-Clause）', 'Mozilla Foundation／Adobe Inc.', 'npm registry → repo lib/', '讀取 PDF、轉成影像', 'lib/pdf.min.js', 'pkg:npm/pdfjs-dist@3.11.174')]:
    add(A, nm, ver, 'library' if 'lang' not in rel else 'data', lic, sup, '美國', src, use, OK, '與 ocr-app 使用相同檔案（SHA-256 相同）', os.path.join(ROOT, rel), purl)
add(A, 'Noto Sans TC（思源黑體繁中）', 'Google Fonts 目前版', 'font', 'OFL-1.1', 'Google LLC／Adobe Inc.', '美國（發行者）',
    'fonts.googleapis.com / fonts.gstatic.com（只有網路版載入）', '畫面字型', WARN,
    'Noto Sans CJK／思源黑體由 Adobe 與 Google 發行，但依公開資料，中文字形由合作字型廠「常州華文」（中國大陸）參與設計。'
    '離線版已不載入；網路版建議改用電腦內建字型（見下方替代方案）。', purl='pkg:generic/google-fonts/noto-sans-tc')
add(A, 'PingFang TC（蘋方，系統字型，僅備援）', 'macOS／iOS 內建', 'font', 'Apple 系統授權', 'Apple Inc.', '美國（發行者）',
    '不隨附、不下載；只在 Apple 裝置上當備援字型', '畫面字型備援', WARN,
    '依公開資料，蘋方由 Apple 與「常州華文」（中國大陸）合作設計。離線版已從字型清單移除；網路版建議一併移除。')
add(A, 'Microsoft JhengHei（微軟正黑體，系統字型）', 'Windows 內建', 'font', 'Windows 系統授權', 'Microsoft Corporation', '美國',
    '不隨附、不下載', '畫面字型（Windows）', OK)
add(A, 'LiHei Pro（儷黑 Pro，系統字型，僅離線版）', 'macOS 內建', 'font', 'Apple 系統授權', '華康科技 DynaComware（Apple 隨附）', '臺灣',
    '不隨附、不下載', '畫面字型（Mac 備援）', OK)
add(A, 'Google Identity Services（gsi/client）', '線上最新版', 'service', 'Google API 服務條款', 'Google LLC', '美國',
    'accounts.google.com（只有按「Google Drive」時才載入）', 'Google 帳號登入（drive.file 權限）', OK)
add(A, 'Google API Loader＋Google Picker', '線上最新版', 'service', 'Google API 服務條款', 'Google LLC', '美國',
    'apis.google.com（只有按「Google Drive」時才載入）', 'Google Drive 選檔', OK)
add(A, 'OneDrive File Picker SDK（OneDrive.js）', '7.2', 'service', 'Microsoft API 使用條款', 'Microsoft Corporation', '美國',
    'js.live.net（只有按「OneDrive」時才載入）', 'OneDrive 選檔', OK)
add(A, 'GitHub Pages（網站託管）', '—', 'service', 'GitHub 服務條款', 'GitHub, Inc.（Microsoft）', '美國',
    'chaifeng666.github.io', '網路版網站託管（離線版不需要）', OK)
add(A, '政府電子採購網 拒絕往來廠商查詢（外部連結）', '—', 'service', '政府網站', '行政院公共工程委員會', '臺灣',
    'web.pcc.gov.tw（只有按「查拒往」時開新分頁）', '查詢賣方是否為拒絕往來廠商', OK, '按鈕以網址參數 suppNo 帶入統編，開新分頁顯示查詢結果；只傳送統編（公開資料），與手動查詢相同。')
add(A, '離線啟動檔 start-windows.bat／server.ps1／start-mac-linux.command', '—', 'application', '專案自有', 'chaifeng666', '臺灣',
    'repo offline/', '離線版本機伺服器（只接受 localhost）', OK, '使用系統內建 Windows PowerShell（Microsoft）或 python3（PSF）')

# ---------- B. ocr-app ----------
add(B, 'ocr.html（自行開發）', '2026-10-24', 'application', '專案自有', 'chaifeng666', '臺灣',
    '/mnt/project-files/ocr-app/', '主程式', OK, '', ocr('ocr.html'))
add(B, 'tesseract.js', '5.1.1', 'library', 'Apache-2.0', 'Project Naptha（Kevin Kwok 等）', '美國',
    'npm registry → OCR-offline.zip lib/；線上版從 jsDelivr 下載', 'OCR 主程式（瀏覽器端）', OK, '',
    ocr('lib/tesseract.min.js'), 'pkg:npm/tesseract.js@5.1.1')
add(B, 'tesseract.js worker', '5.1.1', 'library', 'Apache-2.0', 'Project Naptha', '美國',
    '同上', 'OCR 背景執行緒', OK, '', ocr('lib/worker.min.js'), 'pkg:npm/tesseract.js@5.1.1#worker')
add(B, 'tesseract.js-core（Tesseract OCR 5＋Leptonica，WebAssembly）', '5.1.1', 'library',
    'Apache-2.0（Tesseract）／BSD-2-Clause（Leptonica）', 'Tesseract：HP → Google，開源社群維護；Leptonica：Dan Bloomberg', '美國',
    '同上', 'OCR 引擎', OK, '一般版與 SIMD 加速版各一', ocr('lib/core/tesseract-core-simd-lstm.wasm.js'), 'pkg:npm/tesseract.js-core@5.1.1')
add(B, 'Tesseract 語言資料 chi_tra（繁體中文）', '4.0.0 best_int', 'data', 'Apache-2.0', 'Google LLC（tessdata）', '美國',
    'tesseract.js 官方語言資料 → lib/lang/', '繁體中文辨識模型', OK, '', ocr('lib/lang/chi_tra.traineddata.gz'),
    'pkg:generic/tessdata/chi_tra@4.0.0')
add(B, 'Tesseract 語言資料 eng（英文）', '4.0.0 best_int', 'data', 'Apache-2.0', 'Google LLC（tessdata）', '美國',
    '同上', '英文辨識模型', OK, '', ocr('lib/lang/eng.traineddata.gz'), 'pkg:generic/tessdata/eng@4.0.0')
add(B, 'PDF.js（pdfjs-dist）', '3.11.174', 'library', 'Apache-2.0', 'Mozilla Foundation', '美國',
    'npm registry → lib/', '讀取 PDF、轉成影像', OK, '', ocr('lib/pdf.min.js'), 'pkg:npm/pdfjs-dist@3.11.174')
add(B, 'PDF.js worker', '3.11.174', 'library', 'Apache-2.0', 'Mozilla Foundation', '美國',
    '同上', 'PDF 背景解析', OK, '', ocr('lib/pdf.worker.min.js'), 'pkg:npm/pdfjs-dist@3.11.174#worker')
add(B, 'Adobe CMap 字元對照表（pdfjs-dist/cmaps）', '3.11.174 隨附', 'data', 'BSD-3-Clause', 'Adobe Inc.', '美國',
    '同上 → lib/cmaps/', '正確讀出 PDF 內的中日韓文字', OK,
    '檔名有 GB、UniGB 的是「簡體字編碼」的對照表，由 Adobe 製作，不是中國大陸來源的元件；只用來讀取用該編碼存的 PDF。')
add(B, 'SheetJS（xlsx）', '0.18.5', 'library', 'Apache-2.0', 'SheetJS LLC', '美國',
    'npm registry → lib/', '讀寫 Excel', OK, '', ocr('lib/xlsx.full.min.js'), 'pkg:npm/xlsx@0.18.5')
add(B, 'mammoth', '1.8.0', 'library', 'BSD-2-Clause', 'Michael Williamson', '英國',
    'npm registry → lib/', '讀取 Word（.docx）文字', OK, '', ocr('lib/mammoth.browser.min.js'), 'pkg:npm/mammoth@1.8.0')
add(B, 'JSZip', '3.10.1', 'library', 'MIT OR GPL-3.0-or-later', 'Stuart Knightley', '英國',
    'npm registry → lib/', '讀取 PPT、Word 壓縮內容', OK, '', ocr('lib/jszip.min.js'), 'pkg:npm/jszip@3.10.1')
add(B, '手寫數字辨識模型 DIGIT_NET（自行訓練）', '2026-10-13', 'data', '專案自有（訓練資料 CC BY-SA 3.0）',
    'chaifeng666；訓練資料 MNIST（LeCun、Cortes、Burges，源自美國 NIST）', '臺灣（模型）／美國（資料）',
    '內嵌於 ocr.html', '領據的身分證、電話等數字欄', OK, 'int8 權重直接寫在 ocr.html 裡，不需下載')
add(B, 'jsDelivr CDN', '—', 'service', 'jsDelivr 服務條款', 'jsDelivr（Prospect One，波蘭）；節點由 Cloudflare、Fastly（美國）提供', '波蘭／美國',
    'cdn.jsdelivr.net（只有線上載入模式；OCR-offline.zip 不使用）', '下載程式庫與語言資料（只下載，不上傳）', OK,
    '網頁的安全限制（CSP）只允許連到 jsDelivr 與本機。')
add(B, '系統字型：Microsoft JhengHei／PingFang TC／Noto Sans TC', '系統內建', 'font', '系統授權', 'Microsoft／Apple／Google', '美國（發行者）',
    '不隨附、不下載（只寫在字型清單裡）', '畫面字型', WARN,
    'PingFang TC 與 Noto Sans TC 依公開資料由「常州華文」（中國大陸）參與中文字形設計；建議從字型清單移除，只留微軟正黑體。')
add(B, '離線啟動檔 start-windows.bat／server.ps1／start-mac-linux.command', '—', 'application', '專案自有', 'chaifeng666', '臺灣',
    'OCR-offline.zip', '本機伺服器（只接受 localhost）', OK, '', ocr('server.ps1'))
add(B, 'PaddleOCR PP-OCRv5', '—', 'library', 'Apache-2.0', '百度（Baidu）', '中國大陸',
    '已於 2026-10-13 移除', '（舊版 OCR 引擎）', GONE, '依使用者規定移除，改用 Tesseract（Google）。目前程式中已無任何引用。')
add(B, 'OpenCC', '—', 'library', 'Apache-2.0', 'BYVoid', '中國大陸',
    '已於 2026-10-13 移除', '（舊版簡繁轉換）', GONE, '依使用者規定移除。目前程式中已無任何引用。')

# ---------- CycloneDX 1.5 ----------
now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0).isoformat()
bom = {'bomFormat': 'CycloneDX', 'specVersion': '1.5', 'serialNumber': f'urn:uuid:{uuid.uuid4()}', 'version': 1,
       'metadata': {'timestamp': now, 'component': {'type': 'application', 'name': 'OCR 輸入資料辨識（ocr-classify-ui＋ocr-app）', 'bom-ref': 'root'},
                    'properties': [{'name': 'policy', 'value': '不使用中國大陸來源的元件、模型、字型、CDN'}]},
       'components': [], 'dependencies': []}
TYPE = {'library': 'library', 'application': 'application', 'data': 'data', 'font': 'file', 'service': 'platform'}
for i, c in enumerate(C):
    comp = {'bom-ref': f'c{i+1}', 'type': TYPE[c['type']], 'name': c['name'], 'version': c['version'],
            'supplier': {'name': c['supplier']},
            'licenses': [{'expression': c['license']}] if c['license'][:1].isascii() and '（' not in c['license'] else [{'license': {'name': c['license']}}],
            'properties': [{'name': k, 'value': v} for k, v in [('應用', c['app']), ('來源國', c['country']), ('取得方式', c['source']),
                                                                  ('用途', c['use']), ('檢查結論', c['result']), ('說明', c['note'])] if v]}
    if c['purl']: comp['purl'] = c['purl']
    if c['sha256']: comp['hashes'] = [{'alg': 'SHA-256', 'content': c['sha256']}]
    bom['components'].append(comp)
os.makedirs(os.path.join(ROOT, 'sbom'), exist_ok=True)
with open(os.path.join(ROOT, 'sbom', 'sbom.cdx.json'), 'w', encoding='utf-8') as f:
    json.dump(bom, f, ensure_ascii=False, indent=2)

# ---------- Markdown ----------
cnt = lambda r: sum(1 for c in C if c['result'] == r)
md = [f'# SBOM 軟體物料清單', '', f'產生時間：{now}　格式：CycloneDX 1.5（[sbom.cdx.json](sbom.cdx.json)）', '',
      '檢查目標：**不使用中國大陸來源的元件、模型、字型、CDN。**', '',
      f'結果：{OK} {cnt(OK)} 項、{WARN} {cnt(WARN)} 項、{GONE} {cnt(GONE)} 項（舊版已移除）。'
      '所有隨附的程式庫與模型都不是中國大陸來源；需注意的只有「字型」。', '',
      '## 需注意項目與替代方案', '',
      '| 項目 | 問題 | 替代方案 |', '|---|---|---|',
      '| Noto Sans TC（Google Fonts，網路版） | 發行者為 Google／Adobe（美國），但中文字形由常州華文（中國大陸）參與設計 | 改用電腦內建字型：微軟正黑體（Windows）、儷黑 Pro（Mac，臺灣華康）。離線版已採用。 |',
      '| PingFang TC（系統備援字型） | 不隨附，但依公開資料由常州華文參與設計 | 從字型清單移除。離線版已移除。 |', '']
for app in (A, B):
    md += [f'## {app}', '', '| # | 元件 | 版本 | 授權 | 開發者／組織 | 來源國 | 取得方式 | 結論 |', '|---|---|---|---|---|---|---|---|']
    for i, c in enumerate([c for c in C if c['app'] == app], 1):
        md.append(f"| {i} | {c['name']} | {c['version']} | {c['license']} | {c['supplier']} | {c['country']} | {c['source']} | {c['result']} |")
    md.append('')
md += ['## 說明', ''] + [f"- **{c['name']}**：{c['note']}" for c in C if c['note']] + ['',
       '來源國以「開發者／發行組織所在地」判斷；SHA-256 雜湊值見 sbom.cdx.json，可用來核對隨附檔案沒有被替換。', '']
with open(os.path.join(ROOT, 'sbom', 'SBOM.md'), 'w', encoding='utf-8') as f:
    f.write('\n'.join(md))

# ---------- Excel ----------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
wb = Workbook(); ws = wb.active; ws.title = 'SBOM'
head = ['#', '應用', '元件', '版本', '類型', '授權', '開發者／組織', '來源國', '取得方式', '用途', '檢查結論', '說明', 'SHA-256', 'PURL']
ws.append(head)
for i, c in enumerate(C, 1):
    ws.append([i, c['app'], c['name'], c['version'], c['type'], c['license'], c['supplier'], c['country'], c['source'], c['use'],
               c['result'], c['note'], c['sha256'], c['purl']])
thin = Side(style='thin', color='C9CED6')
fill = {OK: 'E6F4EA', WARN: 'FFF4CE', GONE: 'EEEEEE'}
for cell in ws[1]:
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F2A5C')
    cell.alignment = Alignment(vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=2):
    r = row[10].value
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True); cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    row[10].fill = PatternFill('solid', fgColor=fill[r])
for col, w in zip('ABCDEFGHIJKLMN', [4, 22, 34, 14, 10, 20, 28, 12, 32, 22, 11, 50, 20, 26]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
s2 = wb.create_sheet('摘要')
for r in [['SBOM 軟體物料清單'], ['產生時間', now], ['檢查目標', '不使用中國大陸來源的元件、模型、字型、CDN'],
          ['通過', cnt(OK)], ['需注意', cnt(WARN)], ['已移除（舊版）', cnt(GONE)], [],
          ['需注意項目', '替代方案'],
          ['Noto Sans TC（網路版字型）', '改用電腦內建：微軟正黑體（Windows）、儷黑 Pro（Mac，臺灣華康）；離線版已採用'],
          ['PingFang TC（系統備援字型）', '從字型清單移除；離線版已移除']]:
    s2.append(r)
s2['A1'].font = Font(bold=True, size=14); s2.column_dimensions['A'].width = 28; s2.column_dimensions['B'].width = 70
wb.save(XLSX_OUT)
print(len(C), 'components ->', XLSX_OUT)
